"""
Synchronous in-process report generation helpers.
"""
import csv
import json
import os
import re
from datetime import date, timedelta

from django.conf import settings
from django.db.models import Count, Q
from django.utils import timezone

from crimes.models import Crime, CrimeCategory
from .models import Report

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


CONTENT_TYPES = {
    'pdf': 'application/pdf',
    'excel': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'csv': 'text/csv',
    'json': 'application/json',
}


def content_type_for_format(fmt):
    return CONTENT_TYPES.get((fmt or 'json').lower(), 'application/octet-stream')


def _default_date_range(params=None):
    params = params or {}
    end = date.today()
    start = end - timedelta(days=30)
    if params.get('start_date'):
        raw = params['start_date']
        start = raw if isinstance(raw, date) else date.fromisoformat(str(raw)[:10])
    if params.get('end_date'):
        raw = params['end_date']
        end = raw if isinstance(raw, date) else date.fromisoformat(str(raw)[:10])
    return start, end


def _normalize_category_key(value):
    """Normalize labels like THEFT_OF_VEHICLE / theft-of-vehicle for loose matching."""
    text = str(value or '').strip().lower()
    text = text.replace('_', ' ').replace('-', ' ')
    text = re.sub(r'\s+', ' ', text)
    return text


def _keys_equivalent(a, b):
    if a == b:
        return True
    # rough singular/plural: "theft of vehicle" ≈ "theft of vehicles"
    if a.rstrip('s') == b.rstrip('s'):
        return True
    return False


def resolve_crime_category_ids(raw_values):
    """
    Resolve a list of category PKs and/or name labels to CrimeCategory IDs.
    Accepts ints, digit strings, display names, and UPPER_SNAKE labels.
    """
    if not raw_values:
        return []

    values = list(raw_values) if not isinstance(raw_values, (str, bytes)) else [raw_values]
    ids = set()
    name_candidates = []

    for value in values:
        if value is None or value == '':
            continue
        if isinstance(value, int) or (isinstance(value, str) and str(value).isdigit()):
            ids.add(int(value))
            continue
        name_candidates.append(str(value).strip())

    if name_candidates:
        name_q = Q()
        for name in name_candidates:
            name_q |= Q(name__iexact=name)
        matched_ids = set(
            CrimeCategory.objects.filter(name_q).values_list('id', flat=True)
        )
        ids.update(matched_ids)

        unmatched = [
            c for c in name_candidates
            if not CrimeCategory.objects.filter(name__iexact=c).exists()
        ]
        if unmatched:
            wanted = {_normalize_category_key(c) for c in unmatched}
            alias_map = {
                'theft': 'stealing',
                'assault': 'offenses',
                'burglary': 'breakings',
                'vandalism': 'criminal damage',
                'drug': 'dangerous drugs',
                'drug offenses': 'dangerous drugs',
                'fraud': 'economic crimes',
                'other': 'other offenses',
                'traffic': 'traffic offenses',
                'economic': 'economic crimes',
            }
            expanded = set(wanted)
            for key in list(wanted):
                if key in alias_map:
                    expanded.add(alias_map[key])
            for cat in CrimeCategory.objects.all().only('id', 'name'):
                cat_key = _normalize_category_key(cat.name)
                if any(_keys_equivalent(cat_key, w) for w in expanded):
                    ids.add(cat.id)

    return sorted(ids)


def _build_summary(report):
    """Aggregate crime stats for the report filters."""
    qs = Crime.objects.filter(
        date__gte=report.start_date,
        date__lte=report.end_date,
        location__isnull=False,
    )
    crime_types = list(report.crime_types.values_list('id', flat=True))
    if crime_types:
        qs = qs.filter(category_id__in=crime_types)
    districts = list(report.districts.values_list('id', flat=True))
    if districts:
        qs = qs.filter(district_id__in=districts)
    neighborhoods = list(report.neighborhoods.values_list('id', flat=True))
    if neighborhoods:
        qs = qs.filter(neighborhood_id__in=neighborhoods)

    if report.center_point and report.radius:
        from django.contrib.gis.measure import D
        qs = qs.filter(location__distance_lte=(report.center_point, D(km=float(report.radius))))

    total = qs.count()
    violent = qs.filter(is_violent=True).count()
    by_category = list(
        qs.values('category__name')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    by_status = list(
        qs.values('status')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    by_district = list(
        qs.values('district__name')
        .annotate(count=Count('id'))
        .order_by('-count')[:20]
    )
    sample = list(
        qs.order_by('-date')[:100].values(
            'id', 'case_number', 'date', 'description', 'category__name',
            'block_address', 'is_violent', 'status',
        )
    )
    return {
        'title': report.title,
        'report_type': report.report_type,
        'description': report.description or '',
        'date_range': {
            'start_date': report.start_date.isoformat(),
            'end_date': report.end_date.isoformat(),
        },
        'totals': {
            'total_crimes': total,
            'violent_crimes': violent,
            'non_violent_crimes': total - violent,
        },
        'by_category': [
            {'name': row['category__name'] or 'Unknown', 'count': row['count']}
            for row in by_category
        ],
        'by_status': [
            {'status': row['status'] or 'unknown', 'count': row['count']}
            for row in by_status
        ],
        'by_district': [
            {'name': row['district__name'] or 'Unknown', 'count': row['count']}
            for row in by_district
        ],
        'sample_crimes': sample,
        'generated_at': timezone.now().isoformat(),
    }


def _ensure_reports_dir():
    reports_dir = os.path.join(settings.MEDIA_ROOT, 'reports')
    os.makedirs(reports_dir, exist_ok=True)
    return reports_dir


def _write_json(path, summary):
    with open(path, 'w', encoding='utf-8') as fh:
        json.dump(summary, fh, indent=2, default=str)


def _write_csv(path, report, summary):
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        writer = csv.writer(fh)
        writer.writerow(['metric', 'value'])
        writer.writerow(['title', report.title])
        writer.writerow(['report_type', report.report_type])
        writer.writerow(['start_date', report.start_date.isoformat()])
        writer.writerow(['end_date', report.end_date.isoformat()])
        writer.writerow(['total_crimes', summary['totals']['total_crimes']])
        writer.writerow(['violent_crimes', summary['totals']['violent_crimes']])
        writer.writerow(['non_violent_crimes', summary['totals']['non_violent_crimes']])
        writer.writerow(['generated_at', summary['generated_at']])
        writer.writerow([])
        writer.writerow(['category', 'count'])
        for row in summary['by_category']:
            writer.writerow([row['name'], row['count']])
        writer.writerow([])
        writer.writerow(['district', 'count'])
        for row in summary['by_district']:
            writer.writerow([row['name'], row['count']])
        writer.writerow([])
        writer.writerow([
            'id', 'case_number', 'date', 'category', 'description',
            'block_address', 'is_violent', 'status',
        ])
        for crime in summary['sample_crimes']:
            writer.writerow([
                crime.get('id'),
                crime.get('case_number'),
                crime.get('date'),
                crime.get('category__name'),
                crime.get('description'),
                crime.get('block_address'),
                crime.get('is_violent'),
                crime.get('status'),
            ])


def _write_excel(path, report, summary):
    if not HAS_OPENPYXL:
        raise RuntimeError('openpyxl is required to generate Excel reports')

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = 'Summary'
    bold = Font(bold=True)

    summary_ws.append(['Title', report.title])
    summary_ws.append(['Report Type', report.report_type])
    summary_ws.append(['Start Date', report.start_date.isoformat()])
    summary_ws.append(['End Date', report.end_date.isoformat()])
    summary_ws.append(['Total Crimes', summary['totals']['total_crimes']])
    summary_ws.append(['Violent Crimes', summary['totals']['violent_crimes']])
    summary_ws.append(['Non-Violent Crimes', summary['totals']['non_violent_crimes']])
    summary_ws.append(['Generated At', summary['generated_at']])
    for cell in summary_ws['A']:
        cell.font = bold

    cat_ws = wb.create_sheet('By Category')
    cat_ws.append(['Category', 'Count'])
    cat_ws['A1'].font = bold
    cat_ws['B1'].font = bold
    for row in summary['by_category']:
        cat_ws.append([row['name'], row['count']])

    dist_ws = wb.create_sheet('By District')
    dist_ws.append(['District', 'Count'])
    dist_ws['A1'].font = bold
    dist_ws['B1'].font = bold
    for row in summary['by_district']:
        dist_ws.append([row['name'], row['count']])

    sample_ws = wb.create_sheet('Sample Crimes')
    headers = [
        'ID', 'Case Number', 'Date', 'Category', 'Description',
        'Address', 'Violent', 'Status',
    ]
    sample_ws.append(headers)
    for cell in sample_ws[1]:
        cell.font = bold
    for crime in summary['sample_crimes']:
        sample_ws.append([
            crime.get('id'),
            crime.get('case_number'),
            str(crime.get('date') or ''),
            crime.get('category__name'),
            crime.get('description'),
            crime.get('block_address'),
            crime.get('is_violent'),
            crime.get('status'),
        ])

    wb.save(path)


def _write_pdf(path, report, summary):
    if not HAS_REPORTLAB:
        raise RuntimeError('reportlab is required to generate PDF reports')

    doc = SimpleDocTemplate(path, pagesize=letter)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        spaceAfter=12,
    )
    body = []

    body.append(Paragraph(report.title or 'Crime Report', title_style))
    body.append(Paragraph(f"Type: {report.get_report_type_display()}", styles['Normal']))
    body.append(Paragraph(
        f"Date range: {report.start_date.isoformat()} to {report.end_date.isoformat()}",
        styles['Normal'],
    ))
    if report.description:
        body.append(Spacer(1, 0.15 * inch))
        body.append(Paragraph(report.description, styles['Normal']))
    body.append(Spacer(1, 0.25 * inch))

    totals = summary['totals']
    stats_data = [
        ['Metric', 'Value'],
        ['Total crimes', str(totals['total_crimes'])],
        ['Violent crimes', str(totals['violent_crimes'])],
        ['Non-violent crimes', str(totals['non_violent_crimes'])],
        ['Generated at', summary['generated_at']],
    ]
    stats_table = Table(stats_data, colWidths=[2.5 * inch, 4 * inch])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a5f')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    body.append(stats_table)
    body.append(Spacer(1, 0.3 * inch))

    body.append(Paragraph('Crime type summary', styles['Heading2']))
    cat_rows = [['Category', 'Count']]
    for row in summary['by_category'][:40]:
        cat_rows.append([row['name'], str(row['count'])])
    if len(cat_rows) == 1:
        cat_rows.append(['(no matching crimes)', '0'])
    cat_table = Table(cat_rows, colWidths=[4.5 * inch, 2 * inch])
    cat_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    body.append(cat_table)

    if summary['by_district']:
        body.append(Spacer(1, 0.25 * inch))
        body.append(Paragraph('Top districts', styles['Heading2']))
        dist_rows = [['District', 'Count']]
        for row in summary['by_district'][:15]:
            dist_rows.append([row['name'], str(row['count'])])
        dist_table = Table(dist_rows, colWidths=[4.5 * inch, 2 * inch])
        dist_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        body.append(dist_table)

    doc.build(body)


def _write_report_file(report, summary):
    reports_dir = _ensure_reports_dir()
    fmt = (report.format or 'json').lower()

    if fmt == 'csv':
        filename = f"report_{report.id}.csv"
        path = os.path.join(reports_dir, filename)
        _write_csv(path, report, summary)
    elif fmt == 'excel':
        filename = f"report_{report.id}.xlsx"
        path = os.path.join(reports_dir, filename)
        _write_excel(path, report, summary)
    elif fmt == 'pdf':
        filename = f"report_{report.id}.pdf"
        path = os.path.join(reports_dir, filename)
        _write_pdf(path, report, summary)
    else:
        filename = f"report_{report.id}.json"
        path = os.path.join(reports_dir, filename)
        _write_json(path, summary)

    return path


def generate_report_sync(report):
    """
    Synchronously generate report content, write a file, and mark completed.
    Returns the updated Report instance.
    """
    report.status = 'processing'
    report.progress = 10
    report.save(update_fields=['status', 'progress', 'updated_at'])

    try:
        summary = _build_summary(report)
        report.progress = 70
        report.save(update_fields=['progress', 'updated_at'])

        path = _write_report_file(report, summary)
        report.file_path = path
        report.file_size = os.path.getsize(path)
        report.status = 'completed'
        report.progress = 100
        report.completed_at = timezone.now()
        report.error_message = None
        report.save()
        return report
    except Exception as exc:
        report.status = 'failed'
        report.error_message = str(exc)
        report.progress = 0
        report.save(update_fields=['status', 'error_message', 'progress', 'updated_at'])
        raise


def create_report_from_params(user, params, template=None):
    """
    Create a Report from request/template params and generate it synchronously.
    """
    params = dict(params or {})
    template_data = (template.template_data if template else {}) or {}
    merged = {**template_data, **params}

    start_date, end_date = _default_date_range(merged)
    title = merged.get('title') or (template.name if template else 'Crime Report')
    description = merged.get('description') or (template.description if template else '')
    report_type = merged.get('report_type') or (template.report_type if template else 'crime_summary')
    fmt = merged.get('format') or 'json'

    report = Report.objects.create(
        user=user,
        title=title,
        description=description or '',
        report_type=report_type,
        start_date=start_date,
        end_date=end_date,
        include_charts=bool(merged.get('include_charts', True)),
        include_map=bool(merged.get('include_map', True)),
        format=fmt if fmt in dict(Report.FORMAT_CHOICES) else 'json',
        status='pending',
        radius=merged.get('radius'),
    )

    crime_types = merged.get('crime_types') or merged.get('crimeTypes') or []
    if crime_types:
        ids = resolve_crime_category_ids(crime_types)
        if ids:
            report.crime_types.set(ids)

    return generate_report_sync(report)
